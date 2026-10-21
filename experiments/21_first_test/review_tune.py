# 맞춤용 9건(화면 10개) 검토 자료: 0.2초마다 사진(위험한계·관절·세훈 판단·판단기 3개 표시) + 엑셀 한 파일
# 시험용 10건은 만들지 않는다. 결과: OmniGuard-data/세훈_작업/맞춤용_검토/
import os, sys, json, numpy as np, cv2
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
HERE = os.path.dirname(os.path.abspath(__file__)); src = open(os.path.join(HERE, "evaluate.py"), encoding="utf-8").read(); exec(src[:src.index("res = {}")])
sys.path.insert(0, os.path.join(X, "09_red_fit")); from common import KP_MIN, to_frame
C = DATA + "collected/"; OUT = DATA + "세훈_작업/맞춤용_검토/"; os.makedirs(OUT, exist_ok=True)
SRC = {"A09_pipe": (C + "A/A09.mp4"), "A10_pipe_view2": C + "A/A10.mp4", "B01_lathe": C + "B/B01.mp4", "B02_lathe_chuck": C + "B/B02.mp4",
       **{f"D{i:02d}": C + f"D/D{i:02d}.mp4" for i in range(1, 14)}}
TUNE = [n for n in sorted(os.listdir(LAB)) if os.path.isdir(LAB + n) and n not in TEST]
JN = {"J13": "J13 원칙", "RULE4": "규칙 4", "QUIET": "조용한 안"}
ST = ["정상", "주의", "정지"]; COL = {0: (90, 90, 90), 1: (0, 200, 255), 2: (0, 0, 255), -1: (40, 40, 40)}  # BGR
XL = {0: "D9D9D9", 1: "FFD966", 2: "FF7C80"}
FB = ImageFont.truetype("C:/Windows/Fonts/malgunbd.ttf", 22)
wb = Workbook(); ws0 = wb.active; ws0.title = "요약"
ws0.append(["영상", "세훈 첫 정지(초)", "닿음(초)"] + [f"{JN[j]} 첫 정지" for j in JN] + [f"{JN[j]} 정지 성공" for j in JN] + ["사진 폴더"])
for name in TUNE:
    P = dict(np.load(POSE + f"{name}_pose.npz", allow_pickle=True)); lab, tc, hs = labels(name); zone, objects, floor, fronts = geometry(name)
    F = J.features(P, zone, floor, objects); t = F["t"]; outs = {j: J.run(j, P, F, fronts) for j in JN}
    s3 = {j: np.where(o == 3, 2, np.where(o >= 1, 1, 0)) for j, o in outs.items()}  # 판단기 감속은 주의로
    win = (t >= tc - 0.5) & (t <= tc); first = {j: next((round(float(t[i]), 1) for i in range(len(t)) if outs[j][i] == 3), None) for j in JN}
    d = OUT + name + "/"; os.makedirs(d, exist_ok=True); ws = wb.create_sheet(name[:28])
    ws.append(["시각(초)", "사진", "세훈 판단", "닿음", "J13 원칙", "규칙 4", "조용한 안", "다른 곳", "세훈 원래 메모", "왜 위험/안전한가 (여기에 적기)"])
    memo = {}
    import csv
    for r in list(csv.reader(open(LAB + f"{name}/판단표.csv", encoding="utf-8-sig")))[1:]:
        memo[round(float(r[0]), 1)] = (r[4] if name not in RULES["칸보정"] else "") if len(r) > 4 else ""
    cap = cv2.VideoCapture(SRC[name])
    for k in sorted(lab):
        i = int(np.argmin(np.abs(t - k))); cap.set(cv2.CAP_PROP_POS_MSEC, float(t[i]) * 1000); ok, img = cap.read()
        if not ok: continue
        hl = F["hull_at"](i); hl = hl if isinstance(hl, list) else [hl]
        ov = img.copy()
        for hh in hl: cv2.fillPoly(ov, [hh.astype(np.int32)], (0, 200, 0))
        img = cv2.addWeighted(ov, 0.25, img, 0.75, 0)
        for hh in hl: cv2.polylines(img, [hh.astype(np.int32)], True, (0, 220, 0), 2)
        for a, b in fronts:
            q = to_frame(P, i, [a, b]).astype(int); cv2.line(img, tuple(q[0]), tuple(q[1]), (255, 0, 255), 3)
        rr = max(3, img.shape[1] // 250)
        for kk, s in zip(P["kps"][i], P["sc"][i]):
            for j in range(17):
                if s[j] >= KP_MIN: cv2.circle(img, (int(kk[j][0]), int(kk[j][1])), rr + (2 if j in (9, 10) else 0), (0, 255, 0) if j in (9, 10) else (255, 220, 0), -1)
        h, w = img.shape[:2]; S = 960 / w; img = cv2.resize(img, (960, int(h * S)))
        y = lab[k]; vals = [("세훈", y)] + [(JN[j], int(s3[j][i])) for j in JN]
        bar = np.zeros((44, 960, 3), np.uint8); cw = 960 // 4
        for n_, (lb, v) in enumerate(vals): cv2.rectangle(bar, (n_ * cw, 0), ((n_ + 1) * cw - 2, 43), COL[v], -1)
        pil = Image.fromarray(cv2.cvtColor(np.vstack([bar, img]), cv2.COLOR_BGR2RGB)); dr = ImageDraw.Draw(pil)
        for n_, (lb, v) in enumerate(vals): dr.text((n_ * cw + 8, 8), f"{lb}: {ST[v]}", font=FB, fill=(255, 255, 255))
        dr.text((8, 50), f"{name}  {k:.1f}초" + ("   ← 닿음" if tc is not None and abs(k - tc) < 0.05 else ""), font=FB, fill=(255, 255, 0))
        fn = f"{k:06.1f}s.jpg"; cv2.imencode(".jpg", cv2.cvtColor(np.array(pil), cv2.COLOR_RGB2BGR), [cv2.IMWRITE_JPEG_QUALITY, 85])[1].tofile(d + fn)
        jv = [int(s3[j][i]) for j in JN]; diff = ", ".join(JN[j] for j, v in zip(JN, jv) if (v == 2) != (y == 2))
        ws.append([k, fn, ST[y], "Y" if tc is not None and abs(k - tc) < 0.05 else "", *[ST[v] for v in jv], diff, memo.get(k, ""), ""])
        row = ws.max_row
        for c_, v in zip("CEFG", [y] + jv): ws[f"{c_}{row}"].fill = PatternFill("solid", fgColor=XL[v])
        if diff: ws[f"H{row}"].font = Font(bold=True, color="C00000")
    for c_, wd in zip("ABCDEFGHIJ", [9, 13, 10, 6, 10, 10, 10, 22, 40, 50]): ws.column_dimensions[c_].width = wd
    ws.freeze_panes = "A2"
    hit = {j: bool(win.any() and (outs[j][win] == 3).all()) for j in JN}
    ws0.append([name, hs, tc] + [first[j] for j in JN] + ["O" if hit[j] else "X" for j in JN] + [f"맞춤용_검토/{name}/"])
    print(name, len(lab), "frames", first, hit, flush=True)
for c_ in "ABCDEFGHIJ": ws0.column_dimensions[c_].width = 16
ws0.append([]); ws0.append(["색: 회색 정상 / 노랑 주의(판단기 감속 포함) / 빨강 정지. '다른 곳' = 세훈과 정지 여부가 다른 판단기"])
wb.save(OUT + "맞춤용_검토.xlsx"); print("saved", OUT + "맞춤용_검토.xlsx")
